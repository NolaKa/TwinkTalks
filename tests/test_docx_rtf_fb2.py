"""Tests for the docx, rtf, and fb2 extractors plus their router wiring."""

import pytest

from twinktalks import extractor


@pytest.fixture
def docx_file(tmp_path):
    from docx import Document
    p = tmp_path / "doc.docx"
    doc = Document()
    doc.add_heading("Chapter 1", level=1)
    doc.add_paragraph("First paragraph of chapter one.")
    doc.add_heading("Section A", level=2)
    doc.add_paragraph("Some content under section A.")
    doc.add_heading("Chapter 2", level=1)
    doc.add_paragraph("Second chapter starts here.")
    doc.save(str(p))
    return p


@pytest.fixture
def rtf_file(tmp_path):
    p = tmp_path / "doc.rtf"
    p.write_text(r"{\rtf1\ansi This is a \b bold\b0 rtf string.}")
    return p


@pytest.fixture
def fb2_file(tmp_path):
    p = tmp_path / "doc.fb2"
    p.write_text(
        "<?xml version='1.0' encoding='utf-8'?>\n"
        "<FictionBook xmlns='http://www.gribuser.ru/xml/fictionbook/2.0'>\n"
        "  <body>\n"
        "    <section><title><p>Prologue</p></title>"
        "      <p>Once upon a time.</p>"
        "      <section><title><p>Sub one</p></title>"
        "        <p>Nested paragraph.</p>"
        "      </section>"
        "    </section>\n"
        "    <section><title><p>Chapter One</p></title>"
        "      <p>Story begins.</p>"
        "    </section>\n"
        "  </body>\n"
        "  <body name='notes'><section><title><p>Footnote</p></title><p>Skip me.</p></section></body>\n"
        "</FictionBook>\n",
        encoding="utf-8",
    )
    return p


class TestDocx:
    def test_text_in_document_order(self, docx_file):
        text = extractor.extract_text(str(docx_file))
        assert "First paragraph of chapter one." in text
        assert text.index("Chapter 1") < text.index("First paragraph")
        assert text.index("Section A") < text.index("Some content")
        assert text.index("Chapter 2") < text.index("Second chapter")

    def test_toc_from_headings(self, docx_file):
        toc = extractor.extract_toc(str(docx_file))
        titles = [(c.title, c.level) for c in toc]
        assert ("Chapter 1", 1) in titles
        assert ("Section A", 2) in titles
        assert ("Chapter 2", 1) in titles

    def test_get_item_count(self, docx_file):
        assert extractor.get_item_count(str(docx_file)) == 1


class TestRtf:
    def test_strips_control_codes(self, rtf_file):
        text = extractor.extract_text(str(rtf_file))
        assert "rtf string" in text
        assert "\\rtf" not in text
        assert "\\b" not in text

    def test_no_toc(self, rtf_file):
        # RTF has no structured headings — empty TOC is the contract.
        assert extractor.extract_toc(str(rtf_file)) == []


class TestFb2:
    def test_text_includes_paragraphs(self, fb2_file):
        text = extractor.extract_text(str(fb2_file))
        assert "Once upon a time." in text
        assert "Story begins." in text
        # The notes body must NOT leak into the main text.
        assert "Skip me." not in text

    def test_toc_walks_sections(self, fb2_file):
        toc = extractor.extract_toc(str(fb2_file))
        titles_levels = [(c.title, c.level) for c in toc]
        assert ("Prologue", 1) in titles_levels
        assert ("Sub one", 2) in titles_levels
        assert ("Chapter One", 1) in titles_levels


class TestRouter:
    def test_supported_extensions_includes_new_formats(self):
        for ext in (".docx", ".rtf", ".fb2"):
            assert ext in extractor.SUPPORTED_EXTENSIONS

    def test_pdf_kwargs_dropped_for_docx(self, docx_file):
        # OCR / page_range are PDF-only; the router must filter them out
        # before calling docx_extractor (which doesn't accept them).
        text = extractor.extract_text(
            str(docx_file), ocr=True, ocr_language="eng", page_range=(1, 5),
        )
        assert "First paragraph" in text
